import json, uuid, io, zipfile, re
from contextlib import asynccontextmanager
from pathlib import Path
from typing import AsyncIterator
import httpx
from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from .config import OLLAMA_URL, DEFAULT_MODEL, WORKSPACE, ROOT
from . import store
from .tools import TOOL_DEFINITIONS, run_tool, ToolError, safe_path

SYSTEM = """You are FRIDAY, a capable, precise local assistant. You run through the user's local Ollama instance. Be honest about capabilities and verification. Use tools only when helpful. Filesystem tools are confined to FRIDAY's private workspace. Do not claim to have controlled the desktop, browsed the web, or completed an action unless a tool actually did it. Before destructive, sensitive, financial, or externally consequential actions, ask for explicit confirmation. Describe plans briefly for multi-step work. Reply in the user's language when possible."""

@asynccontextmanager
async def lifespan(app: FastAPI):
    store.initialize()
    yield

app = FastAPI(title="FRIDAY AI Local API", version="0.1.0", lifespan=lifespan)

class ChatRequest(BaseModel):
    conversation_id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    message: str = Field(min_length=1, max_length=12000)
    model: str | None = None
    approve_workspace_writes: bool = False

class DocumentRequest(BaseModel):
    kind: str = Field(pattern="^(docx|pptx|xlsx|pdf|md)$")
    title: str = Field(min_length=1, max_length=120)
    body: str = Field(min_length=1, max_length=50000)

async def ollama_json(path: str, payload: dict):
    try:
        async with httpx.AsyncClient(timeout=httpx.Timeout(180, connect=3)) as client:
            response = await client.post(f"{OLLAMA_URL}{path}", json=payload)
            response.raise_for_status()
            return response.json()
    except httpx.ConnectError as e:
        raise HTTPException(503, "Ollama is unavailable. Start Ollama locally and try again.") from e
    except httpx.HTTPStatusError as e:
        raise HTTPException(e.response.status_code, f"Ollama error: {e.response.text[:500]}") from e
    except httpx.TimeoutException as e:
        raise HTTPException(504, "Ollama took too long to respond.") from e

@app.get("/api/health")
async def health():
    try:
        async with httpx.AsyncClient(timeout=2) as client:
            r = await client.get(f"{OLLAMA_URL}/api/tags")
            r.raise_for_status()
            models = [x.get("name", "") for x in r.json().get("models", [])]
        return {"status": "ready", "ollama": True, "models": models, "default_model": DEFAULT_MODEL, "workspace": str(WORKSPACE)}
    except Exception as e:
        return {"status": "degraded", "ollama": False, "models": [], "error": str(e), "workspace": str(WORKSPACE)}

@app.get("/api/models")
async def models():
    try:
        async with httpx.AsyncClient(timeout=3) as client:
            r = await client.get(f"{OLLAMA_URL}/api/tags")
            r.raise_for_status()
            return {"models": r.json().get("models", []), "default_model": DEFAULT_MODEL}
    except httpx.ConnectError as e:
        raise HTTPException(503, "Ollama is unavailable. Start Ollama locally and install a model.") from e

async def stream_ollama(model: str, messages: list[dict], with_tools: bool = True) -> AsyncIterator[dict]:
    body = {"model": model, "messages": messages, "stream": True, "options": {"temperature": 0.35}}
    if with_tools: body["tools"] = TOOL_DEFINITIONS
    try:
        async with httpx.AsyncClient(timeout=httpx.Timeout(240, connect=5)) as client:
            async with client.stream("POST", f"{OLLAMA_URL}/api/chat", json=body) as r:
                if r.status_code >= 400:
                    raw = (await r.aread()).decode(errors="replace")
                    raise RuntimeError(f"Ollama error ({r.status_code}): {raw[:500]}")
                async for line in r.aiter_lines():
                    if line:
                        yield json.loads(line)
    except httpx.ConnectError as e:
        raise RuntimeError("Ollama is unavailable. Start Ollama locally and try again.") from e

async def sse(event: str, data: dict):
    return f"event: {event}\ndata: {json.dumps(data, ensure_ascii=False)}\n\n"

@app.post("/api/chat")
async def chat(req: ChatRequest):
    model = req.model or DEFAULT_MODEL
    if not model:
        raise HTTPException(400, "Choose or install an Ollama model first (for example: ollama pull qwen2.5:7b).")
    store.add_message(req.conversation_id, "user", req.message)
    history = store.get_messages(req.conversation_id, 24)
    messages = [{"role": "system", "content": SYSTEM}, *history]
    async def events():
        full = []
        try:
            yield await sse("status", {"state": "thinking", "model": model})
            collected, calls = [], []
            async for chunk in stream_ollama(model, messages, True):
                msg = chunk.get("message", {})
                text = msg.get("content", "")
                if text:
                    collected.append(text)
                    yield await sse("token", {"text": text})
                if msg.get("tool_calls"):
                    calls.extend(msg["tool_calls"])
                if chunk.get("error"):
                    raise RuntimeError(chunk["error"])
            full.extend(collected)
            if calls:
                assistant_msg = {"role": "assistant", "content": "".join(collected), "tool_calls": calls}
                messages.append(assistant_msg)
                results = []
                for call in calls[:6]:
                    fn = call.get("function", {})
                    name = fn.get("name", "")
                    args = fn.get("arguments", {})
                    if isinstance(args, str): args = json.loads(args)
                    yield await sse("tool", {"name": name, "status": "running"})
                    try:
                        result = run_tool(name, args, req.approve_workspace_writes)
                    except (ToolError, OSError, ValueError, TypeError) as e:
                        result = {"error": str(e)}
                    results.append({"role": "tool", "content": json.dumps(result, ensure_ascii=False), "name": name})
                    yield await sse("tool", {"name": name, "status": "done", "result": result})
                messages.extend(results)
                collected2 = []
                async for chunk in stream_ollama(model, messages, False):
                    text = chunk.get("message", {}).get("content", "")
                    if text:
                        collected2.append(text)
                        yield await sse("token", {"text": text})
                    if chunk.get("error"): raise RuntimeError(chunk["error"])
                full.extend(collected2)
            answer = "".join(full).strip()
            if answer: store.add_message(req.conversation_id, "assistant", answer)
            yield await sse("done", {"conversation_id": req.conversation_id, "verified": bool(calls), "model": model})
        except Exception as e:
            yield await sse("error", {"message": str(e)})
    return StreamingResponse(events(), media_type="text/event-stream", headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})

@app.get("/api/conversations")
def conversations():
    return {"items": store.list_conversations()}

@app.get("/api/conversations/{conversation_id}")
def conversation(conversation_id: str):
    return {"conversation_id": conversation_id, "messages": store.get_messages(conversation_id, 100)}

@app.get("/api/files")
def files(path: str = "."):
    p = safe_path(path)
    if not p.exists() or not p.is_dir(): raise HTTPException(404, "Workspace folder not found")
    return {"path": str(p.relative_to(WORKSPACE)), "items": [{"name": x.name, "kind": "folder" if x.is_dir() else "file", "bytes": x.stat().st_size if x.is_file() else None} for x in sorted(p.iterdir(), key=lambda x: (not x.is_dir(), x.name.lower()))[:300] if not x.name.startswith(".")]}

@app.post("/api/documents")
def create_document(req: DocumentRequest):
    dest = WORKSPACE / "generated"
    dest.mkdir(exist_ok=True)
    slug = re.sub(r"[^a-z0-9-]+", "-", req.title.lower()).strip("-")[:50] or "friday-document"
    path = dest / f"{slug}.{req.kind}"
    i = 2
    while path.exists():
        path = dest / f"{slug}-{i}.{req.kind}"; i += 1
    try:
        if req.kind == "md":
            path.write_text(f"# {req.title}\n\n{req.body}\n", encoding="utf-8")
        elif req.kind == "docx":
            from docx import Document
            d = Document(); d.add_heading(req.title, 0)
            for para in req.body.split("\n\n"): d.add_paragraph(para)
            d.save(path)
        elif req.kind == "pptx":
            from pptx import Presentation
            from pptx.util import Inches, Pt
            prs = Presentation(); slides = [x.strip() for x in req.body.split("\n\n") if x.strip()]
            for index, text in enumerate(slides or [req.body]):
                slide = prs.slides.add_slide(prs.slide_layouts[0] if index == 0 else prs.slide_layouts[1])
                slide.shapes.title.text = req.title if index == 0 else text.split("\n")[0][:100]
                if index == 0 and len(slide.placeholders) > 1: slide.placeholders[1].text = "Created with FRIDAY AI"
                elif len(slide.placeholders) > 1: slide.placeholders[1].text = "\n".join(text.split("\n")[1:]) or text
            prs.save(path)
        elif req.kind == "xlsx":
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            wb = Workbook(); ws = wb.active; ws.title = "FRIDAY"
            ws.append([req.title]); ws["A1"].font = Font(bold=True, size=16, color="FFFFFF"); ws["A1"].fill = PatternFill("solid", fgColor="073F52")
            for line in req.body.splitlines(): ws.append(line.split("\t"))
            ws.column_dimensions["A"].width = 32; wb.save(path)
        elif req.kind == "pdf":
            from reportlab.lib.pagesizes import letter
            from reportlab.pdfgen import canvas
            c = canvas.Canvas(str(path), pagesize=letter); width, height = letter
            c.setFont("Helvetica-Bold", 18); c.drawString(54, height-60, req.title[:90]); c.setFont("Helvetica", 11)
            y = height-95
            for line in req.body.splitlines():
                if y < 60: c.showPage(); c.setFont("Helvetica", 11); y = height-60
                c.drawString(54, y, line[:105]); y -= 16
            c.save()
        # Basic artifact verification: file exists and is nonempty.
        if not path.is_file() or path.stat().st_size == 0: raise RuntimeError("Output verification failed")
        return {"created": str(path.relative_to(WORKSPACE)), "bytes": path.stat().st_size, "verified": True}
    except ImportError as e:
        raise HTTPException(501, f"Install the document dependencies to create {req.kind}: {e}") from e

@app.get("/api/download/{relative_path:path}")
def download(relative_path: str):
    p = safe_path(relative_path)
    if not p.is_file(): raise HTTPException(404, "File not found")
    return FileResponse(p, filename=p.name)

@app.get("/api/zip/{relative_path:path}")
def zip_folder(relative_path: str):
    p = safe_path(relative_path)
    if not p.is_dir(): raise HTTPException(404, "Folder not found")
    memory = io.BytesIO()
    with zipfile.ZipFile(memory, "w", zipfile.ZIP_DEFLATED) as archive:
        for item in p.rglob("*"):
            if item.is_file() and not any(part.startswith(".") for part in item.relative_to(p).parts):
                archive.write(item, arcname=str(Path(p.name) / item.relative_to(p)))
    if memory.tell() == 0: raise HTTPException(400, "Folder is empty")
    memory.seek(0)
    return StreamingResponse(memory, media_type="application/zip", headers={"Content-Disposition": f'attachment; filename="{p.name}.zip"'})

frontend = ROOT / "frontend"
if frontend.exists():
    app.mount("/", StaticFiles(directory=frontend, html=True), name="frontend")
